"""Gate: the one-sheet hunting list is one sheet, eight columns, and hunts only. Ryan, 2026-09-12.

The list Ryan uses is `hunting-list.xlsx`, written by src/write_hunting_docs.py from the same row list as the
long documents. This RUNS the writer with HOME pointed at a scratch folder, so the real documents are never
touched, and inspects the file:

  L1  exactly one sheet
  L2  exactly the eight columns, in order: what it is, year, club, men held, what a find would add, newspaper,
      date, link -- and no more
  L3  every row is one of the four hunts: nobody held, almost nobody held, boundary season, surname only
  L4  no row is a missing-field row: an "almost nobody" row holds fewer men than the writer's FEWER_THAN, and
      nothing anywhere on the sheet says "not a hunt"
  L5  a row naming a newspaper carries its link in the row, and its first link is clickable

  python3 src/gate_hunting_list.py [--script PATH]   exit 1 = FAIL
"""
import os, re, sys, subprocess, tempfile

HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
COLUMNS = ["what it is", "year", "club", "men held", "what a find would add", "newspaper", "date", "link"]
KINDS = ("nobody held", "almost nobody held", "boundary season", "surname only: ")
FAILS = []


def check(ok, msg):
    print(("  ok    " if ok else "  FAIL  ") + msg)
    if not ok: FAILS.append(msg)


def main():
    script = sys.argv[sys.argv.index("--script") + 1] if "--script" in sys.argv else os.path.join(HERE, "write_hunting_docs.py")
    home = tempfile.mkdtemp(prefix="gate-hunting-list-", dir=os.environ.get("GATE_SCRATCH") or None)
    docs = os.path.join(home, "Dropbox", "Football Archive", "docs"); os.makedirs(docs)
    os.symlink(os.path.expanduser("~/Dropbox/Football Archive/reports"), os.path.join(home, "Dropbox", "Football Archive", "reports"))
    print(f"writer: {script}\nwriting into scratch: {docs} (the real documents are not touched)")
    p = subprocess.run([sys.executable, script], env={**os.environ, "HOME": home}, capture_output=True, text=True)
    path = os.path.join(docs, "hunting-list.xlsx")
    if p.returncode != 0 or not os.path.exists(path):
        check(False, f"L0 the writer produced hunting-list.xlsx (exit {p.returncode}; {(p.stderr or p.stdout)[-300:]})")
        print(f"\nGATE FAILED ({len(FAILS)})"); return 1
    print(p.stdout.strip().splitlines()[-1])
    from openpyxl import load_workbook
    wb = load_workbook(path)
    check(len(wb.sheetnames) == 1, f"L1 exactly one sheet ({wb.sheetnames})")
    ws = wb[wb.sheetnames[0]]
    head = [c.value for c in ws[1]]
    check(head == COLUMNS, f"L2 exactly the eight columns, in order ({head})")
    body = [r for r in ws.iter_rows(min_row=2) if any(c.value not in (None, "") for c in r)]
    ix = {h: i for i, h in enumerate(head)} if head == COLUMNS else {}
    if not ix:
        print(f"\nGATE FAILED ({len(FAILS)})"); return 1
    bad_kind = [r[ix["what it is"]].value for r in body if not str(r[ix["what it is"]].value or "").startswith(KINDS)]
    check(not bad_kind, f"L3 every row is one of the four hunts ({len(body)} rows; {len(bad_kind)} not: {bad_kind[:4]})")
    sys.path.insert(0, os.path.dirname(os.path.abspath(script)))
    import importlib.util
    spec = importlib.util.spec_from_file_location("writer_under_test", script); W = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(W)
    too_many = [(r[ix["club"]].value, r[ix["men held"]].value) for r in body
                if r[ix["what it is"]].value == "almost nobody held" and not (isinstance(r[ix["men held"]].value, int) and r[ix["men held"]].value < W.FEWER_THAN)]
    marker = [r[0].row for r in body if any("not a hunt" in str(c.value or "").lower() for c in r)]
    check(not too_many and not marker, f"L4 no missing-field row: 'almost nobody' rows under {W.FEWER_THAN} men ({len(too_many)} not: {too_many[:3]}); "
                                       f"'not a hunt' on {len(marker)} rows")
    no_link = [r[ix["club"]].value for r in body if str(r[ix["newspaper"]].value or "").strip() and not str(r[ix["link"]].value or "").strip()]
    no_click = [r[ix["club"]].value for r in body if re.search(r"https?://", str(r[ix["link"]].value or "")) and not r[ix["link"]].hyperlink]
    check(not no_link and not no_click, f"L5 every cited row carries its link in the row ({len(no_link)} do not: {no_link[:3]}), "
                                        f"first link clickable ({len(no_click)} not: {no_click[:3]})")
    print("\nGATE PASSED" if not FAILS else f"\nGATE FAILED ({len(FAILS)})")
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
